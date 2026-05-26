import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "src" / "scheduler"))


def test_default_schedule_request_contains_seed_data():
    from scheduler.service import get_default_schedule_request

    request = get_default_schedule_request()

    assert request["fellow_groups"]["NCC_JR"] == ["NCC Raya", "NCC Joseph"]
    assert request["fellow_groups"]["NCC_SR"] == ["NCC David", "NCC Prash"]
    assert "Stroke Gabi" in request["fellow_groups"]["STROKE"]
    assert "NH Adam" in request["fellow_groups"]["NH"]
    assert request["fellow_groups"]["LIA"] == ["NCC Lia"]
    assert request["fellow_week_pairs"]["NCC Prash"][:3] == [1, 7, 25]
    assert "Clinic/Elective" in request["shifts"]
    assert any(
        rule["kind"] == "vacation_request_policy"
        for rule in request["annual_rules"]["rules"]
    )


def test_build_schedule_workbook_returns_xlsx_bytes():
    from scheduler.service import build_schedule_workbook

    result = {
        "request": {
            "fellow_groups": {
                "NCC_JR": ["NCC Raya"],
                "STROKE": ["Stroke Gabi"],
            },
        },
        "shifts_for_fellows": {
            "NCC Raya": ["MICU"] * 52,
            "Stroke Gabi": ["Stroke"] * 52,
        },
        "fellows_for_shifts": {
            "NCC1": ["NCC Raya"] * 52,
            "NCC2": [""] * 52,
            "Extra": [""] * 52,
            "Swing": [""] * 52,
            "Stroke": ["Stroke Gabi"] * 52,
            "Telestroke/Clinic": [""] * 52,
            "Stroke_Supervisory": [""] * 52,
        },
    }

    workbook = build_schedule_workbook(result)

    assert isinstance(workbook, bytes)
    assert io.BytesIO(workbook).read(2) == b"PK"


def test_per_fellow_workbook_reports_active_soft_constraint_violations_only(monkeypatch, tmp_path):
    import openpyxl

    from scheduler import service
    from scheduler.service import build_schedule_workbook

    standing_config = tmp_path / "standing.yaml"
    standing_config.write_text("""
rules:
  - name: hard_service_profile_is_omitted
    kind: service_profile
    active: true
    strength: hard
    fellow_groups: [NCC_JR]
    totals:
      - shifts: [MICU]
        relation: exactly
        weeks: 50
  - name: inactive_soft_rule_is_omitted
    kind: service_profile
    active: false
    strength: soft
    fellow_groups: [NCC_SR]
    totals:
      - shifts: [MICU]
        relation: exactly
        weeks: 10
  - name: minimize_swing_deficit
    kind: minimize_uncovered_shift_weeks
    active: true
    strength: minimize
    fellow_groups: [NCC_JR, NCC_SR]
    shifts: [Swing]
  - name: ncc_four_week_block_preference
    kind: block_shift_set_choice
    active: true
    strength: soft
    fellow_groups: [NCC_JR]
    block_size: 4
    choices:
      - [NCC1, Swing]
      - [NCC2, Swing]
    allow_none: true
  - name: ncc_jr_service_profile
    kind: service_profile
    active: true
    strength: soft
    fellow_groups: [NCC_JR]
    zero_shifts: [Stroke]
    totals:
      - shifts: [NCC1, NCC2]
        relation: exactly
        weeks: 2
      - shifts: [Elec]
        relation: at_least
        weeks: 2
    window_totals:
      - shifts: [MICU]
        relation: exactly
        weeks: 2
        window: [0, 4]
""")
    monkeypatch.setattr(service, "STANDING_RULE_CONFIG", standing_config)

    junior_schedule = ["Elec"] * 52
    junior_schedule[0] = "MICU"
    junior_schedule[1] = "MICU"
    junior_schedule[4] = "NCC1"
    junior_schedule[5] = "NCC2"
    junior_schedule[6] = "Stroke"
    senior_schedule = ["Elec"] * 52
    senior_schedule[0] = "Swing"

    workbook = build_schedule_workbook({
        "request": {
            "fellow_groups": {
                "NCC_JR": ["NCC Junior"],
                "NCC_SR": ["NCC Senior"],
                "STROKE": ["Stroke Fellow"],
            },
            "fellow_week_pairs": {},
            "annual_rules": {
                "rules": [
                    {
                        "name": "preferred_stroke",
                        "kind": "specific_assignment",
                        "active": True,
                        "fellow": "NCC Senior",
                        "week": 2,
                        "shift": "Stroke",
                        "strength": "soft",
                    },
                    {
                        "name": "backup_group_fulfilled",
                        "kind": "specific_assignment",
                        "active": True,
                        "fellows": ["NCC Junior", "NCC Senior"],
                        "week": 0,
                        "shift": "Swing",
                        "strength": "soft",
                    },
                ],
            },
        },
        "shifts_for_fellows": {
            "NCC Junior": junior_schedule,
            "NCC Senior": senior_schedule,
            "Stroke Fellow": ["Stroke"] * 52,
        },
        "fellows_for_shifts": {},
    })

    wb = openpyxl.load_workbook(io.BytesIO(workbook))
    ws = wb["Per-Fellow Schedule"]
    header_row = [cell.value for cell in ws[1]]
    report_start = header_row.index("Fellow") + 1
    headers = [
        ws.cell(row=1, column=report_start + offset).value
        for offset in range(7)
    ]
    report_rows = [
        tuple(
            ws.cell(row=row, column=report_start + offset).value
            for offset in range(7)
        )
        for row in range(2, ws.max_row + 1)
        if any(
            ws.cell(row=row, column=report_start + offset).value is not None
            for offset in range(7)
        )
    ]

    assert headers == ["Fellow", "Group", "Rule", "Scope", "Requirement", "Current", "Violation"]
    assert (
        "NCC Junior",
        "NCC_JR",
        "ncc_jr_service_profile",
        "Zero",
        "0 weeks of Stroke",
        1,
        1,
    ) in report_rows
    assert (
        "NCC Junior",
        "NCC_JR",
        "ncc_four_week_block_preference",
        "Weeks 4-7",
        "4-week block must match NCC1/Swing or NCC2/Swing or no NCC1/NCC2/Swing",
        "NCC1, NCC2, Stroke, Elec",
        1,
    ) in report_rows
    assert (
        None,
        None,
        "minimize_swing_deficit",
        "All weeks",
        "minimize uncovered weeks of Swing",
        51,
        51,
    ) in report_rows
    assert (
        "NCC Senior",
        "NCC_SR",
        "preferred_stroke",
        "Week 2",
        "NCC Senior assigned to Stroke",
        "Elec",
        1,
    ) in report_rows
    assert all(row[6] > 0 for row in report_rows)
    assert all(row[2] != "hard_service_profile_is_omitted" for row in report_rows)
    assert all(row[2] != "inactive_soft_rule_is_omitted" for row in report_rows)
    assert all(row[2] != "backup_group_fulfilled" for row in report_rows)
    assert all(row[4] != "exactly 2 weeks of NCC1/NCC2" for row in report_rows)
    assert all(row[4] != "exactly 2 weeks of MICU" for row in report_rows)
    assert all(row[0] != "Stroke Fellow" for row in report_rows)


def test_flask_default_schedule_endpoint_returns_seed_data():
    from scheduler import app

    client = app.test_client()
    response = client.get("/api/default-schedule")

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["fellow_groups"]["NCC_JR"] == ["NCC Raya", "NCC Joseph"]
    assert payload["fellow_groups"]["LIA"] == ["NCC Lia"]


def test_solve_schedule_normalizes_and_returns_optimizer_result(monkeypatch):
    from scheduler import service

    captured = {}

    def fake_optimize_schedule(**kwargs):
        captured.update(kwargs)
        return (
            {"NCC Raya": ["MICU"] * 52},
            {"NCC1": ["NCC Raya"] * 52, "NCC2": [""] * 52},
        )

    monkeypatch.setattr(service, "optimize_schedule", fake_optimize_schedule)

    result = service.solve_schedule({
        "fellow_groups": {"NEW_GROUP": ["NCC Raya"]},
        "fellow_week_pairs": {"NCC Raya": [0, 1]},
    })

    assert captured["fellow_groups"] == {"NEW_GROUP": ["NCC Raya"]}
    assert captured["fellow_week_pairs"] == {"NCC Raya": [0, 1]}
    assert result["shifts_for_fellows"]["NCC Raya"][0] == "MICU"


def test_solve_schedule_loads_standing_and_annual_constraints(monkeypatch, tmp_path):
    from scheduler import service

    standing_config = tmp_path / "standing.yaml"
    standing_config.write_text("""
rules:
  - name: ncc_coverage
    kind: ncc_coverage
    active: true
    strength: hard
    fellow_groups: [NCC_JR]
    shifts: [NCC1, NCC2, Swing]
    swing_deficit: 8
""")
    monkeypatch.setattr(service, "STANDING_RULE_CONFIG", standing_config)
    captured = {}

    def fake_optimize_schedule(**kwargs):
        captured.update(kwargs)
        return (
            {"NCC Raya": ["MICU"] * 52},
            {"NCC1": ["NCC Raya"] * 52, "NCC2": [""] * 52},
        )

    monkeypatch.setattr(service, "optimize_schedule", fake_optimize_schedule)

    service.solve_schedule({
        "fellow_groups": {"NCC_JR": ["NCC Raya"]},
        "fellow_week_pairs": {"NCC Raya": [0, 1, 2]},
        "annual_rules": {
            "rules": [
                {
                    "name": "vacation_requests",
                    "kind": "vacation_request_policy",
                    "active": True,
                    "hard_request_count": 2,
                },
                {
                    "name": "first_week",
                    "kind": "specific_assignment",
                    "active": True,
                    "fellow_groups": ["NCC_JR"],
                    "shift": "Stroke",
                    "week": 1,
                    "strength": "soft",
                },
            ],
        },
    })

    kinds = [constraint.kind for constraint in captured["constraints"]]
    assert kinds == [
        "ncc_coverage",
        "specific_assignment",
        "specific_assignment",
        "specific_assignment",
        "specific_assignment",
    ]
    assert captured["constraints"][0].params["swing_deficit"] == 8
    assert captured["constraints"][2].strength.value == "hard"
    assert captured["constraints"][-1].strength.value == "soft"


def test_normalize_schedule_request_does_not_backfill_default_fellows():
    from scheduler.service import normalize_schedule_request

    request = normalize_schedule_request({
        "fellow_groups": {"NEW_GROUP": ["NCC New"]},
        "fellow_week_pairs": {},
    })

    assert request["fellow_groups"] == {"NEW_GROUP": ["NCC New"]}
    assert request["fellow_week_pairs"] == {}
    assert "NCC Prash" not in request["fellow_week_pairs"]


def test_normalize_schedule_request_rejects_requests_for_missing_fellows():
    import pytest

    from scheduler.service import normalize_schedule_request

    with pytest.raises(ValueError, match="Vacation request references unknown fellow: NCC Prash"):
        normalize_schedule_request({
            "fellow_groups": {"NEW_GROUP": ["NCC New"]},
            "fellow_week_pairs": {"NCC Prash": [1]},
        })


def test_normalize_schedule_request_rejects_duplicate_fellows_across_groups():
    import pytest

    from scheduler.service import normalize_schedule_request

    with pytest.raises(ValueError, match="Fellow NCC New appears in multiple fellow_groups"):
        normalize_schedule_request({
            "fellow_groups": {
                "GROUP_A": ["NCC New"],
                "GROUP_B": ["NCC New"],
            },
        })
