from __future__ import annotations

import pytest
from scheduler.schedule_import import parse_schedule_file, ImportResult


CSV_SIMPLE = b"""NCC Alaric,NCC Bertram,Stroke Gabi
MICU,NCC1,Stroke
MICU,NCC1,Telestroke/Clinic
NCC1,MICU,Stroke
NCC1,MICU,Elec
"""

CSV_WITH_WEEK_COL = b"""Week,NCC Alaric,NCC Bertram,Weekend NCC1,Night Mon
1,MICU,NCC1,NCC Alaric,NCC Bertram
2,MICU,NCC1,NCC Bertram,NCC Alaric
3,NCC1,MICU,NCC Alaric,NCC Bertram
"""


def test_parse_csv_basic():
    result = parse_schedule_file(CSV_SIMPLE, "schedule.csv")
    assert result.fellow_names == ["NCC Alaric", "NCC Bertram", "Stroke Gabi"]
    assert result.num_weeks == 4
    assert "MICU" in result.shifts_found
    assert "NCC1" in result.shifts_found
    assert "Stroke" in result.shifts_found
    assert "Telestroke/Clinic" in result.shifts_found
    assert "Elec" in result.shifts_found
    assert result.assignments["NCC Alaric"] == ["MICU", "MICU", "NCC1", "NCC1"]
    assert result.assignments["NCC Bertram"] == ["NCC1", "NCC1", "MICU", "MICU"]


def test_parse_csv_skips_weekend_night_columns():
    result = parse_schedule_file(CSV_WITH_WEEK_COL, "schedule.csv")
    assert result.fellow_names == ["NCC Alaric", "NCC Bertram"]
    assert "Weekend NCC1" not in result.fellow_names
    assert "Night Mon" not in result.fellow_names
    assert result.num_weeks == 3


def test_parse_csv_empty_raises():
    with pytest.raises(ValueError, match="Empty CSV"):
        parse_schedule_file(b"", "schedule.csv")


def test_parse_csv_no_fellows_raises():
    with pytest.raises(ValueError, match="No fellow columns"):
        parse_schedule_file(b"Week,Weekend NCC1\n1,Someone\n", "schedule.csv")


def test_parse_csv_with_blanks():
    csv_data = b"Fellow A,Fellow B\nMICU,NCC1\n,NCC1\nMICU,\n"
    result = parse_schedule_file(csv_data, "test.csv")
    assert result.assignments["Fellow A"] == ["MICU", "", "MICU"]
    assert result.assignments["Fellow B"] == ["NCC1", "NCC1", ""]
    assert "" not in result.shifts_found


def test_parse_xlsx():
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["NCC Alaric", "NCC Bertram"])
    ws.append(["MICU", "NCC1"])
    ws.append(["NCC1", "MICU"])

    buf = BytesIO()
    wb.save(buf)

    result = parse_schedule_file(buf.getvalue(), "schedule.xlsx")
    assert result.fellow_names == ["NCC Alaric", "NCC Bertram"]
    assert result.num_weeks == 2
    assert result.assignments["NCC Alaric"] == ["MICU", "NCC1"]


# ---------------------------------------------------------------------------
# Workbook (3-column-per-fellow) format tests
# ---------------------------------------------------------------------------

def _make_workbook_bytes(fellow_names, data_rows, *, merge_headers=True):
    """Helper to create workbook-format XLSX bytes.

    fellow_names: list of fellow names
    data_rows: list of tuples, each with (week_num, *values)
               where values are triplets (weekday, weekend, nights) per fellow
    merge_headers: if True, merge fellow name cells across 3 columns
    """
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active

    # Row 1: "Week" + fellow names (each spanning 3 columns)
    row1 = ["Week"]
    for name in fellow_names:
        row1.extend([name, None, None])
    ws.append(row1)

    # Merge fellow name cells if requested
    if merge_headers:
        for i, name in enumerate(fellow_names):
            start_col = 2 + i * 3  # 1-indexed: B, E, H, ...
            end_col = start_col + 2
            ws.merge_cells(start_row=1, start_column=start_col,
                           end_row=1, end_column=end_col)

    # Row 2: "Week" + repeating "Weekday", "Weekend", "Nights"
    row2 = ["Week"]
    for _ in fellow_names:
        row2.extend(["Weekday", "Weekend", "Nights"])
    ws.append(row2)

    # Data rows
    for dr in data_rows:
        ws.append(list(dr))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_workbook_format_detected():
    """Workbook format is detected when row 2 contains 'Weekday'."""
    xlsx_bytes = _make_workbook_bytes(
        ["Alice", "Bob"],
        [
            (1, "MICU", "NCC1", "Mon", "NCC1", "Stroke", "Tue"),
            (2, "NCC1", "Stroke", "Wed", "MICU", "NCC2", "Thu"),
        ],
    )
    result = parse_schedule_file(xlsx_bytes, "workbook.xlsx")
    # Should parse successfully as workbook format
    assert result.fellow_names == ["Alice", "Bob"]
    assert result.num_weeks == 2


def test_workbook_format_extracts_weekday_only():
    """Only the Weekday column values are imported, not Weekend or Nights."""
    xlsx_bytes = _make_workbook_bytes(
        ["Alice", "Bob"],
        [
            (1, "MICU", "NCC1", "Mon", "NCC1", "Stroke", "Tue"),
            (2, "Elec", "NCC2", "Wed", "MICU", "NCC1", "Thu"),
            (3, "NCC1", "Stroke", "Fri", "Elec", "NCC2", "Sat"),
        ],
    )
    result = parse_schedule_file(xlsx_bytes, "workbook.xlsx")
    assert result.assignments["Alice"] == ["MICU", "Elec", "NCC1"]
    assert result.assignments["Bob"] == ["NCC1", "MICU", "Elec"]
    # Weekend/Night values should NOT appear in shifts_found
    assert "Mon" not in result.shifts_found
    assert "Tue" not in result.shifts_found
    assert "Wed" not in result.shifts_found
    # Weekday values SHOULD appear
    assert "MICU" in result.shifts_found
    assert "Elec" in result.shifts_found
    assert "NCC1" in result.shifts_found


def test_workbook_format_fellow_names_from_merged_cells():
    """Fellow names are correctly extracted even with merged cells (None padding)."""
    xlsx_bytes = _make_workbook_bytes(
        ["Alice", "Bob", "Charlie"],
        [
            (1, "MICU", "X", "Y", "NCC1", "X", "Y", "Elec", "X", "Y"),
        ],
        merge_headers=True,
    )
    result = parse_schedule_file(xlsx_bytes, "workbook.xlsx")
    assert result.fellow_names == ["Alice", "Bob", "Charlie"]
    assert result.num_weeks == 1


def test_workbook_format_unmerged_headers():
    """Fellow names work when headers are not merged (name only in first col of group)."""
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    # Row 1 without merging — name only in first position, rest None
    ws.append(["Week", "Alice", None, None, "Bob", None, None])
    ws.append(["Week", "Weekday", "Weekend", "Nights", "Weekday", "Weekend", "Nights"])
    ws.append([1, "MICU", "NCC1", "Mon", "NCC1", "Stroke", "Tue"])

    buf = BytesIO()
    wb.save(buf)

    result = parse_schedule_file(buf.getvalue(), "workbook.xlsx")
    assert result.fellow_names == ["Alice", "Bob"]
    assert result.assignments["Alice"] == ["MICU"]
    assert result.assignments["Bob"] == ["NCC1"]


def test_workbook_format_with_empty_weekday_cells():
    """Empty weekday cells are handled as empty strings."""
    xlsx_bytes = _make_workbook_bytes(
        ["Alice"],
        [
            (1, "MICU", "NCC1", "Mon"),
            (2, "", "NCC2", "Tue"),
            (3, "Elec", "", "Wed"),
        ],
    )
    result = parse_schedule_file(xlsx_bytes, "workbook.xlsx")
    assert result.assignments["Alice"] == ["MICU", "", "Elec"]
    assert "" not in result.shifts_found


def test_simple_xlsx_format_still_works():
    """Backward compatibility: simple format (no 'Weekday' in row 2) still works."""
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Alice", "Bob"])
    ws.append(["MICU", "NCC1"])
    ws.append(["NCC1", "MICU"])
    ws.append(["Elec", "NCC1"])

    buf = BytesIO()
    wb.save(buf)

    result = parse_schedule_file(buf.getvalue(), "schedule.xlsx")
    assert result.fellow_names == ["Alice", "Bob"]
    assert result.num_weeks == 3
    assert result.assignments["Alice"] == ["MICU", "NCC1", "Elec"]
    assert result.assignments["Bob"] == ["NCC1", "MICU", "NCC1"]
