from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
import yaml

try:
    from . import annual_rules, standing_rules
    from .date_to_week_index import date_to_week_index
    from .main import W, optimize_schedule
    from .semantic_constraints import ConstraintStrength
except ImportError:  # pragma: no cover - supports running from src/scheduler
    import annual_rules
    import standing_rules
    from date_to_week_index import date_to_week_index
    from main import W, optimize_schedule
    from semantic_constraints import ConstraintStrength

STANDING_RULE_CONFIG = Path(__file__).resolve().parents[2] / "config" / "standing" / "stanford-fellowship.yaml"


DEFAULT_SHIFTS = [
    "SICU",
    "MICU",
    "NS",
    "Anaesthesia",
    "NCC1",
    "NCC2",
    "Swing",
    "Elec",
    "Vac",
    "Stroke",
    "Telestroke/Clinic",
    "Clinic/Elective",
    "SCVMC Rehab",
    "NIR",
    "ISC",
]


def get_default_schedule_request() -> Dict[str, Any]:
    return {
        "fellow_groups": {
            "NCC_JR": ["NCC Raya", "NCC Joseph"],
            "NCC_SR": ["NCC David", "NCC Prash"],
            "STROKE": ["Stroke Gabi", "Stroke Jeff", "Stroke Victoria", "Stroke Parshva"],
            "CCM": [
                "CCM Ariana",
                "CCM Bert",
                "CCM Chloe",
                "CCM Dennis",
                "CCM Edwina",
                "CCM Frank",
                "CCM George",
                "CCM Helen",
                "CCM Iago",
                "CCM Jake",
                "CCM Kyle",
                "CCM Liana",
                "CCM Mary",
                "CCM Ning",
                "CCM Oyo",
            ],
            "NH": ["NH Adam"],
            "LIA": ["NCC Lia"],
        },
        "shifts": list(DEFAULT_SHIFTS),
        "fellow_week_pairs": {
            "NCC Prash": [
                date_to_week_index((2025, 7, 10)),
                date_to_week_index((2025, 8, 20)),
                date_to_week_index((2025, 12, 25)),
                date_to_week_index((2026, 3, 30)),
                date_to_week_index((2026, 4, 6)),
                date_to_week_index((2026, 4, 13)),
                date_to_week_index((2026, 4, 20)),
            ],
            "NCC David": [
                date_to_week_index((2025, 8, 6)),
                date_to_week_index((2025, 8, 13)),
                date_to_week_index((2025, 12, 31)),
            ],
            "NCC Raya": [
                date_to_week_index((2025, 11, 25)),
                date_to_week_index((2026, 3, 18)),
            ],
            "NCC Joseph": [
                date_to_week_index((2025, 12, 25)),
                date_to_week_index((2026, 3, 9)),
                date_to_week_index((2025, 10, 20)),
                date_to_week_index((2026, 5, 18)),
                date_to_week_index((2025, 9, 16)),
            ],
            "Stroke Gabi": [
                date_to_week_index((2025, 10, 28)),
                date_to_week_index((2026, 3, 21)),
                date_to_week_index((2026, 5, 5)),
                date_to_week_index((2025, 9, 16)),
            ],
            "Stroke Jeff": [
                date_to_week_index((2025, 12, 30)),
                date_to_week_index((2026, 3, 27)),
                date_to_week_index((2025, 10, 12)),
            ],
            "Stroke Victoria": [
                date_to_week_index((2025, 10, 10)),
                date_to_week_index((2026, 5, 5)),
                date_to_week_index((2025, 9, 30)),
                date_to_week_index((2025, 9, 16)),
            ],
            "Stroke Parshva": [
                date_to_week_index((2026, 5, 5)),
                date_to_week_index((2025, 10, 14)),
                date_to_week_index((2025, 12, 30)),
                date_to_week_index((2025, 9, 16)),
            ],
        },
        "annual_rules": _default_annual_rules(),
    }


def normalize_schedule_request(raw_request: Dict[str, Any] | None) -> Dict[str, Any]:
    if raw_request is None:
        raise ValueError("Request body must be YAML or JSON.")

    if not isinstance(raw_request, dict):
        raise ValueError("Schedule request must be an object.")

    request: Dict[str, Any] = {
        "fellow_groups": _fellow_groups(raw_request.get("fellow_groups", {})),
        "shifts": _string_list(raw_request.get("shifts", DEFAULT_SHIFTS), "shifts"),
    }
    request["fellow_week_pairs"] = _week_pairs(raw_request.get("fellow_week_pairs", {}))
    if "annual_rules" in raw_request:
        request["annual_rules"] = raw_request["annual_rules"]
    _validate_request_references(request)

    return request


def solve_schedule(raw_request: Dict[str, Any] | None) -> Dict[str, Any]:
    request = normalize_schedule_request(raw_request)
    shifts_for_fellows, fellows_for_shifts = optimize_schedule(
        fellow_groups=request["fellow_groups"],
        shifts=request["shifts"],
        fellow_week_pairs=request["fellow_week_pairs"],
        constraints=_configured_constraints(request),
    )
    return {
        "request": request,
        "shifts_for_fellows": shifts_for_fellows,
        "fellows_for_shifts": fellows_for_shifts,
    }


def build_schedule_workbook(result: Dict[str, Any]) -> bytes:
    request = result["request"]
    shifts_for_fellows = result["shifts_for_fellows"]
    fellows_for_shifts = result["fellows_for_shifts"]

    wb = openpyxl.Workbook()
    _build_per_fellow_sheet(wb.active, request, shifts_for_fellows)
    ws = wb.create_sheet("NCC+Stroke Shift Schedule")
    _build_per_shift_sheet(ws, fellows_for_shifts)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _string_list(value: Any, key: str) -> List[str]:
    if not isinstance(value, list) or any(not isinstance(item, str) for item in value):
        raise ValueError(f"{key} must be a list of strings.")
    return [item.strip() for item in value if item.strip()]


def _fellow_groups(value: Any) -> Dict[str, List[str]]:
    if not isinstance(value, dict):
        raise ValueError("fellow_groups must be an object.")

    groups: Dict[str, List[str]] = {}
    for group_name, fellows in value.items():
        if not isinstance(group_name, str) or not group_name.strip():
            raise ValueError("fellow_groups keys must be non-empty strings.")
        groups[group_name.strip()] = _string_list(fellows, f"fellow_groups.{group_name}")
    return groups


def _week_pairs(value: Any) -> Dict[str, List[int]]:
    if not isinstance(value, dict):
        raise ValueError("fellow_week_pairs must be an object.")

    pairs: Dict[str, List[int]] = {}
    for fellow, weeks in value.items():
        if not isinstance(fellow, str) or not isinstance(weeks, list):
            raise ValueError("fellow_week_pairs must map fellow names to week lists.")
        parsed_weeks = []
        for week in weeks:
            if not isinstance(week, int) or week < 0 or week >= W:
                raise ValueError("Vacation weeks must be integers from 0 through 51.")
            parsed_weeks.append(week)
        pairs[fellow] = parsed_weeks
    return pairs


def _validate_request_references(request: Dict[str, Any]) -> None:
    known_fellows: set[str] = set()
    for fellows in request["fellow_groups"].values():
        for fellow in fellows:
            if fellow in known_fellows:
                raise ValueError(f"Fellow {fellow} appears in multiple fellow_groups")
            known_fellows.add(fellow)

    for fellow in request["fellow_week_pairs"]:
        if fellow not in known_fellows:
            raise ValueError(f"Vacation request references unknown fellow: {fellow}")


def _configured_constraints(request: Dict[str, Any]):
    standing_config = _read_yaml_mapping(STANDING_RULE_CONFIG)
    return [
        *standing_rules.constraints_from_config(standing_config),
        *annual_rules.constraints_from_config(
            request.get("annual_rules"),
            fellow_week_pairs=request.get("fellow_week_pairs", {}),
        ),
    ]


def _read_yaml_mapping(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text())
    if not isinstance(data, dict):
        raise ValueError(f"{path} must contain a YAML mapping.")
    return data


def _default_annual_rules() -> Dict[str, Any]:
    return {
        "rules": [
            {
                "name": "vacation_requests",
                "kind": "vacation_request_policy",
                "active": True,
                "hard_request_count": 3,
            },
            {
                "name": "jeff_abpn_stroke",
                "kind": "specific_assignment",
                "active": True,
                "fellow": "Stroke Jeff",
                "week": 11,
                "shift": "Stroke",
                "strength": "hard",
                "reason": "ABPN",
            },
            {
                "name": "abpn_backup_telestroke",
                "kind": "specific_assignment",
                "active": True,
                "fellows": ["NCC Prash", "NCC David"],
                "week": 11,
                "shift": "Telestroke/Clinic",
                "strength": "soft",
                "reason": "ABPN backup coverage",
            },
            {
                "name": "first_week_senior_on_stroke",
                "kind": "specific_assignment",
                "active": True,
                "fellow_groups": ["NCC_SR"],
                "week": 1,
                "shift": "Stroke",
                "strength": "hard",
            },
            {
                "name": "first_week_telestroke",
                "kind": "specific_assignment",
                "active": True,
                "fellow_groups": ["NCC_JR", "NCC_SR", "STROKE"],
                "week": 1,
                "shift": "Telestroke/Clinic",
                "strength": "hard",
            },
            {
                "name": "isc",
                "kind": "isc",
                "active": True,
                "fellow_groups": ["STROKE"],
                "date": [2026, 2, 5],
                "strength": "hard",
            },
            {
                "name": "fourth_block_two_micu_fellows",
                "kind": "fourth_block_two_micu_fellows",
                "active": True,
                "fellow_groups": ["NCC_JR", "NCC_SR"],
                "strength": "hard",
            },
            {
                "name": "nh_no_first_block_ncc",
                "kind": "stroke_no_block_one_ncc",
                "active": True,
                "fellow_groups": ["NH"],
                "strength": "hard",
            },
            _nh_service_profile_rule(),
        ],
    }


def _nh_service_profile_rule() -> Dict[str, Any]:
    return {
        "name": "nh_service_profile",
        "kind": "service_profile",
        "active": True,
        "fellow_groups": ["NH"],
        "strength": "hard",
        "zero_shifts": [
            "MICU",
            "Anaesthesia",
            "Elec",
            "Vac",
            "NS",
            "SICU",
            "SCVMC Rehab",
            "NIR",
            "Clinic/Elective",
            "ISC",
        ],
        "totals": [
            {"shifts": ["NCC1", "NCC2"], "relation": "exactly", "weeks": 4},
            {"shifts": ["Swing"], "relation": "exactly", "weeks": 1},
            {"shifts": ["Telestroke/Clinic"], "relation": "exactly", "weeks": 3},
            {"shifts": ["Stroke"], "relation": "exactly", "weeks": 4},
        ],
    }


def _build_per_fellow_sheet(ws, request: Dict[str, Any], shifts_for_fellows: Dict[str, List[str]]) -> None:
    ws.title = "Per-Fellow Schedule"
    fellows = [
        fellow
        for group_fellows in request["fellow_groups"].values()
        for fellow in group_fellows
    ]
    _write_table(ws, ["Week"] + fellows, [
        [week] + [shifts_for_fellows.get(fellow, [""] * W)[week] for fellow in fellows]
        for week in range(W)
    ])
    report_start_column = len(fellows) + 4
    _write_soft_constraint_violations(ws, report_start_column, request, shifts_for_fellows)


def _build_per_shift_sheet(ws, fellows_for_shifts: Dict[str, List[str]]) -> None:
    shifts = ["NCC1", "NCC2", "Extra", "Swing", "Stroke", "Telestroke/Clinic", "Stroke_Supervisory"]
    _write_table(ws, ["Week"] + shifts, [
        [week] + [_cell_value(fellows_for_shifts.get(shift, [""] * W)[week]) for shift in shifts]
        for week in range(W)
    ])


def _write_table(ws, headers: List[str], rows: List[List[Any]]) -> None:
    fills = _fills()
    for column, header in enumerate(headers, start=1):
        ws.cell(row=1, column=column, value=header)
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=column, value=_cell_value(value))
            if isinstance(value, str) and value in fills:
                cell.fill = fills[value]
    ws.freeze_panes = "B2"


def _write_soft_constraint_violations(
    ws,
    start_column: int,
    request: Dict[str, Any],
    shifts_for_fellows: Dict[str, List[str]],
) -> None:
    headers = ["Fellow", "Group", "Rule", "Scope", "Requirement", "Current", "Violation"]
    rows = _soft_constraint_violation_rows(request, shifts_for_fellows)
    if not rows:
        return

    for offset, header in enumerate(headers):
        ws.cell(row=1, column=start_column + offset, value=header)
    for row_index, row in enumerate(rows, start=2):
        for offset, value in enumerate(row):
            ws.cell(row=row_index, column=start_column + offset, value=value)


def _soft_constraint_violation_rows(
    request: Dict[str, Any],
    shifts_for_fellows: Dict[str, List[str]],
) -> List[List[Any]]:
    rows: List[List[Any]] = []
    for constraint in _configured_constraints(request):
        if constraint.strength not in {ConstraintStrength.SOFT, ConstraintStrength.MINIMIZE}:
            continue
        if constraint.kind == "service_profile":
            rows.extend(_service_profile_violation_rows(request, shifts_for_fellows, constraint))
        elif constraint.kind == "block_shift_set_choice":
            rows.extend(_block_shift_set_choice_violation_rows(request, shifts_for_fellows, constraint))
        elif constraint.kind == "minimize_uncovered_shift_weeks":
            rows.extend(_minimize_uncovered_shift_weeks_violation_rows(request, shifts_for_fellows, constraint))
        elif constraint.kind == "specific_assignment":
            rows.extend(_specific_assignment_violation_rows(request, shifts_for_fellows, constraint))
    return [row for row in rows if row[-1] > 0]


def _service_profile_violation_rows(
    request: Dict[str, Any],
    shifts_for_fellows: Dict[str, List[str]],
    constraint,
) -> List[List[Any]]:
    rows: List[List[Any]] = []
    rule_name = _constraint_name(constraint)

    for fellow, group in _selected_fellows(request, constraint):
        schedule = shifts_for_fellows.get(fellow, [""] * W)
        for total in constraint.params.get("totals", []):
            shifts = total["shifts"]
            current = _count_shifts(schedule, shifts, 0, W)
            rows.append([
                fellow,
                group,
                rule_name,
                "Total",
                _requirement_text(total["relation"], total["weeks"], shifts),
                current,
                _requirement_gap(total["relation"], total["weeks"], current),
            ])

        for shift in constraint.params.get("zero_shifts", []):
            current = _count_shifts(schedule, [shift], 0, W)
            rows.append([
                fellow,
                group,
                rule_name,
                "Zero",
                f"0 weeks of {shift}",
                current,
                current,
            ])

        for total in constraint.params.get("window_totals", []):
            window_start, window_end = total["window"]
            shifts = total["shifts"]
            current = _count_shifts(schedule, shifts, window_start, window_end)
            rows.append([
                fellow,
                group,
                rule_name,
                f"Weeks {window_start}-{window_end - 1}",
                _requirement_text(total["relation"], total["weeks"], shifts),
                current,
                _requirement_gap(total["relation"], total["weeks"], current),
            ])

        for active_block in constraint.params.get("active_blocks", []):
            rows.extend(_active_block_violation_rows(
                fellow,
                group,
                rule_name,
                schedule,
                active_block,
            ))

    return rows


def _active_block_violation_rows(
    fellow: str,
    group: str,
    rule_name: str,
    schedule: List[str],
    active_block: Dict[str, Any],
) -> List[List[Any]]:
    rows: List[List[Any]] = []
    block_size = active_block["block_size"]
    trigger_shifts = active_block["trigger_shifts"]
    for block_start in range(0, W, block_size):
        block_end = min(block_start + block_size, W)
        if _count_shifts(schedule, trigger_shifts, block_start, block_end) == 0:
            continue
        for count_rule in active_block["counts"]:
            shifts = count_rule["shifts"]
            current = _count_shifts(schedule, shifts, block_start, block_end)
            rows.append([
                fellow,
                group,
                rule_name,
                f"Weeks {block_start}-{block_end - 1}",
                _requirement_text(count_rule["relation"], count_rule["weeks"], shifts),
                current,
                _requirement_gap(count_rule["relation"], count_rule["weeks"], current),
            ])
    return rows


def _block_shift_set_choice_violation_rows(
    request: Dict[str, Any],
    shifts_for_fellows: Dict[str, List[str]],
    constraint,
) -> List[List[Any]]:
    rows: List[List[Any]] = []
    block_size = constraint.params["block_size"]
    choices = constraint.params["choices"]
    trigger_shifts = sorted({shift for choice in choices for shift in choice})
    requirement = _block_shift_set_choice_requirement(block_size, choices, trigger_shifts, constraint)
    for fellow, group in _selected_fellows(request, constraint):
        schedule = shifts_for_fellows.get(fellow, [""] * W)
        for block_start in range(0, W, block_size):
            block_end = min(block_start + block_size, W)
            block_length = block_end - block_start
            choice_matches = [
                _count_shifts(schedule, choice, block_start, block_end) == block_length
                for choice in choices
            ]
            none_matches = (
                constraint.params.get("allow_none", False)
                and _count_shifts(schedule, trigger_shifts, block_start, block_end) == 0
            )
            if any(choice_matches) or none_matches:
                continue
            rows.append([
                fellow,
                group,
                _constraint_name(constraint),
                f"Weeks {block_start}-{block_end - 1}",
                requirement,
                ", ".join(schedule[block_start:block_end]),
                1,
            ])
    return rows


def _block_shift_set_choice_requirement(
    block_size: int,
    choices: List[List[str]],
    trigger_shifts: List[str],
    constraint,
) -> str:
    allowed = [" / ".join(choice).replace(" / ", "/") for choice in choices]
    if constraint.params.get("allow_none", False):
        allowed.append(f"no {'/'.join(trigger_shifts)}")
    return f"{block_size}-week block must match {' or '.join(allowed)}"


def _minimize_uncovered_shift_weeks_violation_rows(
    request: Dict[str, Any],
    shifts_for_fellows: Dict[str, List[str]],
    constraint,
) -> List[List[Any]]:
    selected_fellows = [fellow for fellow, _group in _selected_fellows(request, constraint)]
    shifts = list(constraint.shifts.shifts)
    uncovered = 0
    for week in range(W):
        if not any(
            shifts_for_fellows.get(fellow, [""] * W)[week] in shifts
            for fellow in selected_fellows
        ):
            uncovered += 1
    return [[
        None,
        None,
        _constraint_name(constraint),
        "All weeks",
        f"minimize uncovered weeks of {'/'.join(shifts)}",
        uncovered,
        uncovered,
    ]]


def _specific_assignment_violation_rows(
    request: Dict[str, Any],
    shifts_for_fellows: Dict[str, List[str]],
    constraint,
) -> List[List[Any]]:
    week = constraint.weeks.start
    shift = constraint.shifts.shifts[0]
    selected = _selected_fellows(request, constraint)
    current_by_fellow = [
        (fellow, group, shifts_for_fellows.get(fellow, [""] * W)[week])
        for fellow, group in selected
    ]
    if any(current == shift for _fellow, _group, current in current_by_fellow):
        return []
    if len(current_by_fellow) == 1:
        fellow, group, current = current_by_fellow[0]
        return [[
            fellow,
            group,
            _constraint_name(constraint),
            f"Week {week}",
            f"{fellow} assigned to {shift}",
            current,
            1,
        ]]
    return [[
        None,
        ", ".join(dict.fromkeys(group for _fellow, group, _current in current_by_fellow)),
        _constraint_name(constraint),
        f"Week {week}",
        f"one of {', '.join(fellow for fellow, _group, _current in current_by_fellow)} assigned to {shift}",
        "; ".join(f"{fellow}: {current}" for fellow, _group, current in current_by_fellow),
        1,
    ]]


def _selected_fellows(request: Dict[str, Any], constraint) -> List[tuple[str, str]]:
    group_lookup = _fellow_group_lookup(request)
    if constraint.fellows is None:
        return list(group_lookup.items())
    if constraint.fellows.groups:
        return [
            (fellow, group)
            for group in constraint.fellows.groups
            for fellow in request["fellow_groups"].get(group, [])
        ]
    return [
        (fellow, group_lookup[fellow])
        for fellow in constraint.fellows.names
        if fellow in group_lookup
    ]


def _fellow_group_lookup(request: Dict[str, Any]) -> Dict[str, str]:
    return {
        fellow: group
        for group, fellows in request["fellow_groups"].items()
        for fellow in fellows
    }


def _constraint_name(constraint) -> str:
    return constraint.params.get("name", constraint.kind)


def _count_shifts(schedule: List[str], shifts: List[str], start: int, end: int) -> int:
    shift_set = set(shifts)
    return sum(1 for shift in schedule[start:end] if shift in shift_set)


def _requirement_text(relation: str, weeks: int, shifts: List[str]) -> str:
    relation_text = {
        "at_least": "at least",
        "at_most": "at most",
        "exactly": "exactly",
    }.get(relation, relation)
    return f"{relation_text} {weeks} weeks of {'/'.join(shifts)}"


def _requirement_gap(relation: str, weeks: int, current: int) -> int:
    if relation == "at_least":
        return max(0, weeks - current)
    if relation == "at_most":
        return max(0, current - weeks)
    if relation == "exactly":
        return abs(current - weeks)
    return current - weeks


def _cell_value(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(value)
    return value


def _fills() -> Dict[str, Any]:
    colors = {
        "MICU": "B4C6E7",
        "SICU": "FFE699",
        "NCC1": "C6E0B4",
        "NCC2": "C6E0B4",
        "Swing": "A9D08E",
        "Stroke": "F8CBAD",
        "Telestroke/Clinic": "DCF4D6",
        "NS": "FFF3CC",
        "Anaesthesia": "F8CBAD",
        "Clinic/Elective": "C4D677",
        "NIR": "DDB644",
        "Vac": "D9E2F3",
        "Elec": "E7E6E6",
    }
    return {
        key: openpyxl.styles.PatternFill(start_color=color, end_color=color, fill_type="solid")
        for key, color in colors.items()
    }
