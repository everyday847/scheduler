import json

from flask import Flask, Response, jsonify, request, send_file
from flask_cors import CORS


app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "http://localhost:3000"}})


@app.route("/api/configs", methods=["GET"])
def list_configs():
    """List available annual and standing config files."""
    from .solver_bridge import list_configs
    return jsonify(list_configs())


@app.route("/api/config/<config_type>/<filename>", methods=["GET"])
def get_config(config_type: str, filename: str):
    """Load and return a config file as JSON."""
    from .solver_bridge import CONFIG_DIR
    import yaml
    from datetime import date as date_type

    if config_type not in ("annual", "standing"):
        return jsonify({"error": "config_type must be 'annual' or 'standing'"}), 400

    path = CONFIG_DIR / config_type / filename
    if not path.exists() or not path.suffix == ".yaml":
        return jsonify({"error": f"Config not found: {filename}"}), 404

    data = yaml.safe_load(path.read_text())
    _convert_dates(data)
    return jsonify(data)


def _convert_dates(obj):
    """Recursively convert datetime.date objects to ISO strings."""
    from datetime import date as date_type
    if isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(v, date_type):
                obj[k] = v.isoformat()
            else:
                _convert_dates(v)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            if isinstance(v, date_type):
                obj[i] = v.isoformat()
            else:
                _convert_dates(v)


@app.route("/api/config/annual/<filename>/draft", methods=["GET"])
def get_draft(filename: str):
    """Load draft config (falls back to published)."""
    from .solver_bridge import CONFIG_DIR
    from .draft_persistence import load_draft_or_published, has_draft
    from datetime import date as date_type

    try:
        data = load_draft_or_published(CONFIG_DIR / "annual", filename)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    _convert_dates(data)
    data["_has_draft"] = has_draft(CONFIG_DIR / "annual", filename)
    return jsonify(data)


@app.route("/api/config/annual/<filename>/draft", methods=["PUT"])
def save_draft_endpoint(filename: str):
    """Save draft config."""
    from .solver_bridge import CONFIG_DIR
    from .draft_persistence import save_draft

    body = request.get_json(silent=True)
    if not body:
        return jsonify({"error": "Request body must be JSON"}), 400

    save_draft(CONFIG_DIR / "annual", filename, body)
    return jsonify({"status": "saved"})


@app.route("/api/config/annual/<filename>/draft", methods=["DELETE"])
def discard_draft_endpoint(filename: str):
    """Discard draft config."""
    from .solver_bridge import CONFIG_DIR
    from .draft_persistence import discard_draft

    discard_draft(CONFIG_DIR / "annual", filename)
    return jsonify({"status": "discarded"})


@app.route("/api/config/annual/<filename>/publish", methods=["POST"])
def publish_draft_endpoint(filename: str):
    """Publish draft to replace published config."""
    from .solver_bridge import CONFIG_DIR
    from .draft_persistence import publish_draft

    try:
        publish_draft(CONFIG_DIR / "annual", filename)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    return jsonify({"status": "published"})


@app.route("/api/schedule/import", methods=["POST"])
def schedule_import():
    """Import a CSV/XLSX schedule file and return parsed data."""
    from .schedule_import import parse_schedule_file

    if "file" not in request.files:
        return jsonify({"error": "No file uploaded. Send as multipart with field name 'file'."}), 400

    f = request.files["file"]
    if not f.filename:
        return jsonify({"error": "Empty filename"}), 400

    try:
        result = parse_schedule_file(f.read(), f.filename)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    return jsonify({
        "fellow_names": result.fellow_names,
        "num_weeks": result.num_weeks,
        "shifts_found": sorted(result.shifts_found),
        "assignments": result.assignments,
    })


@app.route("/api/feasibility/check", methods=["POST"])
def feasibility_check():
    """Check feasibility of a rule (solo or paired with another).

    Request: { config, standing_rules, rule, pair_with? }
    Response: { satisfiable: bool, elapsed: float }
    """
    import time
    import yaml
    from .solver_bridge import get_runner, STANDING_RULE_CONFIG
    from .palette_rules import palette_rule_to_constraints
    from .palette_derivations import derive_forbidden_shifts
    from .semantic_constraints import (
        ConstraintLifecycle, ConstraintStrength, FellowSelector, SemanticConstraint,
    )
    from parafrost_scheduler.schedule_types import ScheduleSolverConfig
    from parafrost_scheduler.schedule_encoder import build_full_schedule_opb
    from .night_call_types import NightSolverConfig
    from .weekend_call_types import WeekendSolverConfig

    body = request.get_json(silent=True)
    if not body:
        return jsonify({"error": "Request body required"}), 400

    config_data = body.get("config", {})
    rule = body.get("rule")
    pair_with = body.get("pair_with")
    rules_list = body.get("rules")  # For collective testing (e.g., all shift_totals for a group)

    if not rule and not rules_list:
        return jsonify({"error": "rule or rules is required"}), 400

    try:
        fellow_groups = config_data.get("fellow_groups", {})
        shifts = config_data.get("shifts", [])

        # Build standing constraints
        standing_rules_raw = body.get("standing_rules", [])
        standing_constraints = []
        if standing_rules_raw and isinstance(standing_rules_raw[0], dict) and "type" in standing_rules_raw[0]:
            for sr in standing_rules_raw:
                if not sr.get("active", True):
                    continue
                if sr.get("type") == "full_assignment":
                    standing_constraints.append(SemanticConstraint(
                        kind="full_assignment",
                        lifecycle=ConstraintLifecycle.STANDING_RULE,
                        strength=ConstraintStrength.HARD,
                        fellows=FellowSelector.by_groups(*sr["groups"]),
                        params={"name": sr["name"]},
                    ))
                else:
                    standing_constraints.extend(palette_rule_to_constraints(
                        sr, lifecycle=ConstraintLifecycle.STANDING_RULE,
                    ))
        else:
            from .standing_rules import constraints_from_config as standing_constraints_from_config
            standing_config = yaml.safe_load(STANDING_RULE_CONFIG.read_text())
            standing_constraints = list(standing_constraints_from_config(standing_config))

        # Build probe constraints: standing + target rule(s) + optional pair
        if rules_list:
            probe_rules = rules_list
        else:
            probe_rules = [rule]
        if pair_with:
            probe_rules.append(pair_with)

        probe_constraints = list(standing_constraints)
        for pr in probe_rules:
            if pr.get("active", True):
                probe_constraints.extend(palette_rule_to_constraints(
                    pr, lifecycle=ConstraintLifecycle.ANNUAL_RULE,
                ))

        # Derive forbidden shifts using ALL shift_total rules from the config
        # (not just the probe rules) — otherwise single-rule probes over-restrict
        all_annual_rules = config_data.get("rules", [])
        all_rules_for_derivation = standing_rules_raw + all_annual_rules
        probe_constraints.extend(derive_forbidden_shifts(
            all_rules_for_derivation, shifts, fellow_groups,
        ))

        solver_config = ScheduleSolverConfig(
            fellow_groups=fellow_groups,
            shifts=shifts,
            constraints=probe_constraints,
            night_config=NightSolverConfig(),
            weekend_config=WeekendSolverConfig(),
        )

        t0 = time.time()
        opb, var_map = build_full_schedule_opb(solver_config, soft_bound=None)
        upper = sum(w for _, w in var_map.soft_violations)
        opb_probe, _ = build_full_schedule_opb(solver_config, soft_bound=upper)

        runner = get_runner()
        result = runner.solve(opb_probe, timeout=10.0)
        elapsed = time.time() - t0

        return jsonify({"satisfiable": result.satisfiable, "elapsed": round(elapsed, 2)})
    except Exception as exc:
        return jsonify({"error": str(exc), "satisfiable": False}), 500


@app.route("/api/diagnose/stream", methods=["POST"])
def diagnose_stream():
    """SSE endpoint: streams MUS cores as they're found."""
    body = request.get_json(silent=True)

    def generate():
        from .conflict_diagnosis import extract_mus_cores
        from .solver_bridge import get_runner

        if not body:
            yield _sse("error", {"message": "Request body required"})
            return

        config_data = {
            "fellow_groups": body.get("fellow_groups", {}),
            "shifts": body.get("shifts", []),
            "num_weeks": body.get("num_weeks", 52),
            "rules": body.get("rules", []),
            "locked_assignments": body.get("locked_assignments", {}),
        }
        standing_rules = body.get("standing_rules", [])
        annual_rules = body.get("rules", [])

        try:
            runner = get_runner()
            for event in extract_mus_cores(
                config_data, standing_rules, annual_rules, runner,
                probe_timeout=5.0, max_cores=5,
            ):
                yield _sse(event["type"], event)
        except GeneratorExit:
            pass
        except Exception as exc:
            yield _sse("error", {"message": str(exc)})

    return Response(
        generate(),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/api/schedule/export", methods=["POST"])
def schedule_export():
    """Export current solution as an Excel workbook (3 sheets)."""
    import openpyxl
    from io import BytesIO

    body = request.get_json(silent=True)
    if not body:
        return jsonify({"error": "Request body required"}), 400

    weekly = body.get("weekly_assignments", {})
    weekend = body.get("weekend_assignments", [])
    night = body.get("night_assignments", [])

    wb = openpyxl.Workbook()

    # Sheet 1: Weekly Shifts
    ws = wb.active
    ws.title = "Weekly Shifts"
    fellows = list(weekly.keys())
    num_weeks = max((len(v) for v in weekly.values()), default=0)
    ws.append(["Week"] + fellows)
    for w in range(num_weeks):
        row = [w + 1]
        for f in fellows:
            row.append(weekly[f][w] if w < len(weekly[f]) else "")
        ws.append(row)

    # Sheet 2: Weekend Call
    ws2 = wb.create_sheet("Weekend Call")
    if weekend:
        roles = list(weekend[0].keys())
        ws2.append(["Week"] + roles)
        for w, entry in enumerate(weekend):
            ws2.append([w + 1] + [entry.get(r, "") for r in roles])

    # Sheet 3: Night Call
    ws3 = wb.create_sheet("Night Call")
    if night:
        roles = list(night[0].keys())
        ws3.append(["Day"] + roles)
        for d, entry in enumerate(night):
            ws3.append([d + 1] + [entry.get(r, "") for r in roles])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="schedule.xlsx",
    )


@app.route("/api/schedule/stream", methods=["POST"])
def schedule_stream():
    """SSE endpoint: progressively solves and streams improving schedules."""
    body = request.get_json(silent=True)

    def generate():
        from .solver_bridge import build_solver_config_from_request, get_runner
        from parafrost_scheduler.schedule_optimizer import solve_full_schedule_progressive

        try:
            config = build_solver_config_from_request(body)
            runner = get_runner()
        except Exception as exc:
            yield _sse("error", {"message": str(exc)})
            return

        try:
            for event in solve_full_schedule_progressive(config, runner):
                yield _sse(event["type"], event)
        except GeneratorExit:
            pass
        except Exception as exc:
            yield _sse("error", {"message": str(exc)})

    return Response(
        generate(),
        mimetype="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


def _sse(event_type: str, data: dict) -> str:
    return f"event: {event_type}\ndata: {json.dumps(data)}\n\n"


if __name__ == "__main__":
    app.run(debug=True)
