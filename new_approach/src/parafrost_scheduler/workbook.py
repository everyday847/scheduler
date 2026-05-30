"""Excel workbook generator for the night call schedule.

Produces a two-sheet workbook:
  Sheet 1 "Fellow Schedule" — rows = weeks, groups of 3 columns per fellow
  Sheet 2 "Shift Coverage"  — rows = weeks, columns = shift roles
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from scheduler.call_schedule_common import (
    NIGHT_ROLES,
    WEEKEND_ROLES,
    ParsedCallScheduleCsv,
    is_anaesthesia_service,
    is_clinic_service,
    is_preferred_sunday_following_service,
    parse_call_schedule_csv,
)
from scheduler.night_call_solver import NightScheduleSolution
from scheduler.weekend_call_solver import WeekendScheduleSolution
from scheduler.night_call_solver_policy import (
    CRITERION_ANAESTHESIA,
    CRITERION_CLINIC,
    CRITERION_FRIDAY_WEEKEND_NCC1,
    CRITERION_STROKE,
    CRITERION_SUNDAY_FOLLOWING,
    NightPolicyWeights,
)
from parafrost_scheduler.night_solver_pb import _weeks_with_dual_stroke


# ---------------------------------------------------------------------------
# Colour palettes
# ---------------------------------------------------------------------------

SERVICE_COLORS: dict[str, str] = {
    "MICU": "B4C6E7",
    "MSICU": "B4C6E7",
    "Resifellow MSICU": "B4C6E7",
    "SICU": "FFE699",
    "Elective/SICU": "FFE699",
    "NCC1": "C6E0B4",
    "NCC2": "C6E0B4",
    "Swing": "A9D08E",
    "Stroke": "F8CBAD",
    "Telestroke/Clinic": "DCF4D6",
    "Telestroke": "DCF4D6",
    "Anesthesia": "F8CBAD",
    "Anaesthesia": "F8CBAD",
    "Clinic/Elective": "C4D677",
    "NIR": "DDB644",
    "Vacation": "D9E2F3",
    "NS SCVMC": "FFF3CC",
    "SCVMC": "FFF3CC",
}

FELLOW_CATEGORY_COLORS: dict[str, str] = {
    "ncc_senior": "C6E0B4",
    "ncc_junior": "A9D08E",
    "stroke_senior": "F8CBAD",
    "telestroke": "DCF4D6",
    "ccm": "B4C6E7",
}

FELLOW_CATEGORIES: dict[str, str] = {
    "Cindy Wong": "ncc_senior",
    "Alex Hanson": "ncc_senior",
    "Raya Aliakbar": "ncc_junior",
    "Joseph Conovaloff": "ncc_junior",
    "Aditya Srivatsan": "stroke_senior",
    "Cameron Schmidt": "stroke_senior",
    "Harneet Dhillon": "stroke_senior",
    "Helena Xeros": "stroke_senior",
    "Jinyuan Liu": "telestroke",
    "Sokena Zaidi": "telestroke",
}

_ELECTIVE_GRAY = "E7E6E6"
_RED_FONT = "FF0000"
_THIN = Side(style="thin")
_THICK = Side(style="medium")


# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def _make_fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _bold_font(color: str = "000000") -> Font:
    return Font(bold=True, color=color)


def _normal_font(color: str = "000000") -> Font:
    return Font(color=color)


def _service_color(service: str) -> str | None:
    """Return a hex fill color for a weekday service string, or None."""
    if service in SERVICE_COLORS:
        return SERVICE_COLORS[service]
    # Substring match — try longest key first so "Resifellow MSICU" doesn't
    # accidentally match "MSICU" before being matched as a whole.
    for key in sorted(SERVICE_COLORS, key=len, reverse=True):
        if key in service:
            return SERVICE_COLORS[key]
    # "Elective/..." variants that didn't match anything above
    if "Elective" in service or "elective" in service:
        return _ELECTIVE_GRAY
    return None


def _fellow_category_color(fellow_name: str) -> str | None:
    """Return a hex fill colour for a fellow's shift-coverage cell."""
    if fellow_name in FELLOW_CATEGORIES:
        cat = FELLOW_CATEGORIES[fellow_name]
    elif "CCM" in fellow_name or "Pulmonary" in fellow_name:
        cat = "ccm"
    else:
        return None
    return FELLOW_CATEGORY_COLORS.get(cat)


def _right_border() -> Border:
    return Border(right=_THICK)


def _thin_all_border() -> Border:
    return Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)


# ---------------------------------------------------------------------------
# Violation detection
# ---------------------------------------------------------------------------

def _has_soft_violation(
    parsed: ParsedCallScheduleCsv,
    night_solution: NightScheduleSolution,
    weekend_solution: WeekendScheduleSolution | None,
    week_index: int,
    fellow_name: str,
    hard_criteria: frozenset[str],
    dual_stroke_weeks: frozenset[int] | None = None,
) -> bool:
    """Return True if the fellow has any un-hard-enforced soft violation that week."""
    week_row = parsed.week_rows[week_index]
    weekday_service = week_row.weekday_assignments.get(fellow_name, "")
    night_assignments = night_solution.assignments_by_week[week_index]
    if dual_stroke_weeks is None:
        dual_stroke_weeks = _weeks_with_dual_stroke(parsed)

    for day_of_week, role in enumerate(NIGHT_ROLES):
        if night_assignments.get(role) != fellow_name:
            continue

        is_weekday_night = day_of_week <= 4

        # Anaesthesia criterion: weekday nights only
        if is_weekday_night and CRITERION_ANAESTHESIA not in hard_criteria and is_anaesthesia_service(weekday_service):
            return True

        # Clinic criterion: weekday nights only
        if is_weekday_night and CRITERION_CLINIC not in hard_criteria and is_clinic_service(weekday_service):
            return True

        # Stroke criterion: weekday stroke → weekday nights only (exempt dual-stroke weeks)
        if is_weekday_night and CRITERION_STROKE not in hard_criteria:
            if "Stroke" in weekday_service and "Telestroke" not in weekday_service:
                if week_index not in dual_stroke_weeks:
                    return True

        # Stroke criterion: weekend stroke → weekend nights only (exempt dual-stroke weeks)
        if not is_weekday_night and CRITERION_STROKE not in hard_criteria:
            if weekend_solution is not None:
                wknd_stroke = weekend_solution.assignments_by_week[week_index].get("Weekend Stroke")
                if wknd_stroke == fellow_name and week_index not in dual_stroke_weeks:
                    return True

        # Friday / weekend NCC1 criterion
        if day_of_week == 4 and CRITERION_FRIDAY_WEEKEND_NCC1 not in hard_criteria:
            if weekend_solution is not None:
                wknd = weekend_solution.assignments_by_week[week_index]
                if wknd.get("Weekend NCC1") == fellow_name:
                    return True
            elif week_row.schedule_assignments.get("Weekend NCC1") == fellow_name:
                return True

        # Sunday-following criterion
        if day_of_week == 6 and CRITERION_SUNDAY_FOLLOWING not in hard_criteria:
            if week_index + 1 < len(parsed.week_rows):
                next_service = parsed.week_rows[week_index + 1].weekday_assignments.get(fellow_name, "")
                if not is_preferred_sunday_following_service(next_service):
                    return True

    return False


# ---------------------------------------------------------------------------
# Sheet 1: Fellow Schedule
# ---------------------------------------------------------------------------

_DAY_ABBR = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")


def _night_days_string(night_assignments: dict[str, str], fellow_name: str) -> str:
    days = []
    for day_of_week, role in enumerate(NIGHT_ROLES):
        if night_assignments.get(role) == fellow_name:
            days.append(_DAY_ABBR[day_of_week])
    return ", ".join(days)


def _weekend_role_for_fellow(
    weekend_assignments: dict[str, str],
    fellow_name: str,
) -> str:
    for role in WEEKEND_ROLES:
        if weekend_assignments.get(role) == fellow_name:
            # "Weekend NCC1" → "NCC1", "Weekend Stroke" → "Stroke"
            return role.replace("Weekend ", "")
    return ""


def _build_fellow_schedule_sheet(
    ws,
    parsed: ParsedCallScheduleCsv,
    night_solution: NightScheduleSolution,
    weekend_solution: WeekendScheduleSolution | None,
    hard_criteria: frozenset[str],
) -> None:
    fellows = parsed.fellow_names
    n_fellows = len(fellows)

    # ---- Column widths ----
    # Column A = week number
    ws.column_dimensions["A"].width = 6
    for fi in range(n_fellows):
        base_col = 2 + fi * 3  # 1-based: B=2
        ws.column_dimensions[get_column_letter(base_col)].width = 18     # Weekday
        ws.column_dimensions[get_column_letter(base_col + 1)].width = 8  # Weekend
        ws.column_dimensions[get_column_letter(base_col + 2)].width = 14  # Nights

    # ---- Row 1: merged fellow name headers ----
    ws.cell(row=1, column=1, value="Week").font = _bold_font()
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    for fi, fellow in enumerate(fellows):
        start_col = 2 + fi * 3
        end_col = start_col + 2
        ws.merge_cells(
            start_row=1, start_column=start_col,
            end_row=1, end_column=end_col,
        )
        cell = ws.cell(row=1, column=start_col, value=fellow)
        cell.font = _bold_font()
        cell.alignment = Alignment(horizontal="center")
        # Thick right border after each fellow group
        for col in range(start_col, end_col + 1):
            c = ws.cell(row=1, column=col)
            right_side = _THICK if col == end_col else Side(style=None)
            c.border = Border(
                left=Side(style=None),
                right=right_side,
            )

    # ---- Row 2: sub-headers ----
    ws.cell(row=2, column=1, value="Week").font = _bold_font()
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")
    for fi in range(n_fellows):
        base_col = 2 + fi * 3
        for offset, label in enumerate(("Weekday", "Weekend", "Nights")):
            cell = ws.cell(row=2, column=base_col + offset, value=label)
            cell.font = _bold_font()
            cell.alignment = Alignment(horizontal="center")
        # Thick right border on last column of group
        ws.cell(row=2, column=base_col + 2).border = Border(right=_THICK)

    # ---- Data rows ----
    for week_index, week_row in enumerate(parsed.week_rows):
        row = week_index + 3
        ws.cell(row=row, column=1, value=week_index + 1)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")

        night_assignments = night_solution.assignments_by_week[week_index]
        weekend_assignments = (
            weekend_solution.assignments_by_week[week_index]
            if weekend_solution is not None
            else week_row.schedule_assignments
        )

        for fi, fellow in enumerate(fellows):
            base_col = 2 + fi * 3

            weekday_service = week_row.weekday_assignments.get(fellow, "")
            weekend_role = _weekend_role_for_fellow(weekend_assignments, fellow)
            nights_str = _night_days_string(night_assignments, fellow)

            # Weekday cell
            wd_cell = ws.cell(row=row, column=base_col, value=weekday_service)
            color = _service_color(weekday_service)
            if color:
                wd_cell.fill = _make_fill(color)

            # Weekend cell
            ws.cell(row=row, column=base_col + 1, value=weekend_role)

            # Nights cell
            night_cell = ws.cell(row=row, column=base_col + 2, value=nights_str)
            violation = _has_soft_violation(
                parsed, night_solution, weekend_solution, week_index, fellow, hard_criteria,
            )
            if violation and nights_str:
                night_cell.font = _normal_font(_RED_FONT)

            # Thick right border on last column of group
            ws.cell(row=row, column=base_col + 2).border = Border(right=_THICK)

    # ---- Freeze panes ----
    ws.freeze_panes = ws.cell(row=3, column=2)


# ---------------------------------------------------------------------------
# Sheet 2: Shift Coverage
# ---------------------------------------------------------------------------

_SHIFT_HEADERS = (
    "Week",
    "Wknd NCC1", "Wknd NCC2", "Wknd Stroke",
    "Night Mon", "Night Tue", "Night Wed", "Night Thu",
    "Night Fri", "Night Sat", "Night Sun",
)


def _has_shift_soft_violation(
    parsed: ParsedCallScheduleCsv,
    night_solution: NightScheduleSolution,
    weekend_solution: WeekendScheduleSolution | None,
    week_index: int,
    fellow_name: str,
    col_role: str,  # e.g. "Weekend NCC1" or "Night Fri"
    hard_criteria: frozenset[str],
) -> bool:
    """Check whether this specific shift assignment (fellow, week, role) violates a soft criterion."""
    week_row = parsed.week_rows[week_index]
    weekday_service = week_row.weekday_assignments.get(fellow_name, "")

    if col_role.startswith("Night "):
        # Night assignment — find the day_of_week
        day_abbr = col_role[len("Night "):]
        day_of_week = _DAY_ABBR.index(day_abbr)

        if CRITERION_ANAESTHESIA not in hard_criteria and is_anaesthesia_service(weekday_service):
            return True
        if CRITERION_CLINIC not in hard_criteria and is_clinic_service(weekday_service):
            return True
        if CRITERION_STROKE not in hard_criteria and "Stroke" in weekday_service:
            return True
        if day_of_week == 4 and CRITERION_FRIDAY_WEEKEND_NCC1 not in hard_criteria:
            if weekend_solution is not None:
                wknd = weekend_solution.assignments_by_week[week_index]
                if wknd.get("Weekend NCC1") == fellow_name:
                    return True
            elif week_row.schedule_assignments.get("Weekend NCC1") == fellow_name:
                return True
        if day_of_week == 6 and CRITERION_SUNDAY_FOLLOWING not in hard_criteria:
            if week_index + 1 < len(parsed.week_rows):
                next_service = parsed.week_rows[week_index + 1].weekday_assignments.get(fellow_name, "")
                if not is_preferred_sunday_following_service(next_service):
                    return True

    # Weekend roles don't have direct soft violations (those are in the night solver)
    return False


def _build_shift_coverage_sheet(
    ws,
    parsed: ParsedCallScheduleCsv,
    night_solution: NightScheduleSolution,
    weekend_solution: WeekendScheduleSolution | None,
    hard_criteria: frozenset[str],
) -> None:
    # Column widths
    ws.column_dimensions["A"].width = 6
    for col_idx in range(2, len(_SHIFT_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    # Row 1: headers
    for col_idx, header in enumerate(_SHIFT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _bold_font()
        cell.alignment = Alignment(horizontal="center")

    # Map from display column header → actual role key
    # "Wknd NCC1" → "Weekend NCC1", "Night Mon" → "Night Mon"
    def _col_role(header: str) -> str:
        if header.startswith("Wknd "):
            return "Weekend " + header[len("Wknd "):]
        return header

    # Data rows
    for week_index, week_row in enumerate(parsed.week_rows):
        row = week_index + 2
        ws.cell(row=row, column=1, value=week_index + 1)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")

        night_assignments = night_solution.assignments_by_week[week_index]
        if weekend_solution is not None:
            weekend_assignments = weekend_solution.assignments_by_week[week_index]
        else:
            weekend_assignments = {k: v for k, v in week_row.schedule_assignments.items() if k in WEEKEND_ROLES}

        for col_idx, header in enumerate(_SHIFT_HEADERS[1:], start=2):
            role = _col_role(header)
            if role in WEEKEND_ROLES:
                fellow = weekend_assignments.get(role, "")
            else:
                fellow = night_assignments.get(role, "")

            cell = ws.cell(row=row, column=col_idx, value=fellow)
            if fellow:
                bg = _fellow_category_color(fellow)
                if bg:
                    cell.fill = _make_fill(bg)
                violation = _has_shift_soft_violation(
                    parsed, night_solution, weekend_solution,
                    week_index, fellow, role, hard_criteria,
                )
                if violation:
                    cell.font = _normal_font(_RED_FONT)

    # Freeze panes
    ws.freeze_panes = ws.cell(row=2, column=2)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def write_schedule_workbook(
    parsed: ParsedCallScheduleCsv,
    night_solution: NightScheduleSolution,
    weekend_solution: WeekendScheduleSolution | None,
    output_path: Path,
    *,
    hard_criteria: frozenset[str] = frozenset(),
    weights: NightPolicyWeights = NightPolicyWeights(),
) -> None:
    """Write a color-coded Excel workbook for the combined night + weekend schedule."""
    wb = openpyxl.Workbook()

    # Sheet 1
    ws1 = wb.active
    ws1.title = "Fellow Schedule"
    _build_fellow_schedule_sheet(ws1, parsed, night_solution, weekend_solution, hard_criteria)

    # Sheet 2
    ws2 = wb.create_sheet(title="Shift Coverage")
    _build_shift_coverage_sheet(ws2, parsed, night_solution, weekend_solution, hard_criteria)

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Helper: reconstruct solution objects from saved CSV files
# ---------------------------------------------------------------------------

def parse_solutions_from_csvs(
    night_csv: str | Path,
    weekend_csv: str | Path,
) -> tuple[NightScheduleSolution, WeekendScheduleSolution, ParsedCallScheduleCsv]:
    """Read output CSVs and reconstruct solution + parsed objects.

    The night CSV has columns: [fellow...] [existing_schedule_cols...] [Night Mon ... Night Sun]
    The weekend CSV has columns: [fellow...] [existing_schedule_cols...] [Weekend NCC1 NCC2 Stroke]

    Returns (night_solution, weekend_solution, parsed) where ``parsed`` is built
    from the night CSV (which contains the original weekday assignments).
    """
    night_csv = Path(night_csv)
    weekend_csv = Path(weekend_csv)

    # ---- Parse night CSV ----
    with night_csv.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))

    header = rows[0]
    night_role_set = set(NIGHT_ROLES)
    weekend_role_set = set(WEEKEND_ROLES)

    # Split header into (fellow_names | existing_schedule_cols | night_roles)
    night_role_indices = [i for i, col in enumerate(header) if col in night_role_set]
    weekend_role_indices = [i for i, col in enumerate(header) if col in weekend_role_set]
    fellow_and_sched_indices = [
        i for i in range(len(header))
        if i not in set(night_role_indices) and i not in set(weekend_role_indices)
    ]
    # fellow names are those not in any known schedule column
    fellow_indices = [i for i in fellow_and_sched_indices if header[i] not in weekend_role_set]

    fellow_names = [header[i] for i in fellow_indices]
    existing_schedule_columns_from_header = [header[i] for i in weekend_role_indices]

    # Build ParsedCallScheduleCsv from the night CSV data rows
    from scheduler.call_schedule_common import WeekRow, BLOCK_MARKER
    week_rows: list[WeekRow] = []
    trailing_rows: list[list[str]] = []

    for row in rows[1:]:
        padded = row + [""] * max(0, len(header) - len(row))
        if BLOCK_MARKER in padded:
            trailing_rows.append(padded)
            continue
        weekday_assignments = {header[i]: padded[i] for i in fellow_indices}
        schedule_assignments = {header[i]: padded[i] for i in weekend_role_indices}
        week_rows.append(
            WeekRow(
                weekday_assignments=weekday_assignments,
                schedule_assignments=schedule_assignments,
                raw_row=[padded[i] for i in range(len(fellow_names) + len(existing_schedule_columns_from_header))],
            )
        )

    parsed = ParsedCallScheduleCsv(
        fellow_names=fellow_names,
        existing_schedule_columns=tuple(existing_schedule_columns_from_header),
        week_rows=week_rows,
        trailing_rows=trailing_rows,
    )

    # ---- Reconstruct NightScheduleSolution ----
    assignments_by_week: list[dict[str, str]] = []
    for row in rows[1:]:
        padded = row + [""] * max(0, len(header) - len(row))
        if BLOCK_MARKER in padded:
            continue
        night_assignments = {header[i]: padded[i] for i in night_role_indices}
        assignments_by_week.append(night_assignments)

    night_solution = NightScheduleSolution(assignments_by_week=assignments_by_week)

    # ---- Parse weekend CSV for WeekendScheduleSolution ----
    with weekend_csv.open(newline="", encoding="utf-8-sig") as fh:
        wrows = list(csv.reader(fh))

    wheader = wrows[0]
    w_role_indices = [i for i, col in enumerate(wheader) if col in weekend_role_set]

    weekend_assignments_by_week: list[dict[str, str]] = []
    for row in wrows[1:]:
        padded = row + [""] * max(0, len(wheader) - len(row))
        if BLOCK_MARKER in padded:
            continue
        wa = {wheader[i]: padded[i] for i in w_role_indices}
        weekend_assignments_by_week.append(wa)

    weekend_solution = WeekendScheduleSolution(assignments_by_week=weekend_assignments_by_week)

    return night_solution, weekend_solution, parsed
